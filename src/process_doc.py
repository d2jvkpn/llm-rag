#!/usr/bin/env python3
import os, re

import docx, pptx, tiktoken, ebooklib
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader
from ebooklib import epub
from bs4 import BeautifulSoup


def doc_filename(path):
    filename = os.path.basename(path)

    if len(filename) > 64:
        filename = filename[:61] + "..."

    return filename

# doc_id=md5-xxxxxxxx
def document2chunks(path, doc_id, chunk_size=1000, chunk_overlap=100):
    ext = path.rsplit(".", 1)[-1]
    number_of_pages = 0

    if ext == "pptx":
        return pptx2chunks(path, doc_id)
    elif ext == "pdf":
        return pdf2chunks(path, doc_id, chunk_size, chunk_overlap)
    elif ext == "md":
        return md2chunks(path, doc_id, chunk_size, chunk_overlap)
    elif ext == "epub":
        return epub2chunks(path, doc_id, chunk_size)
    elif ext == "txt":
        with open(path, 'r', encoding='utf-8') as f:
            text = f.read()
        texts = [p.strip() for p in re.split(r'\n\s*\n', text) if p.strip()]
    elif ext == "docx":
        doc = docx.Document(path)
        texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    else:
        raise ValueError("unknown filetype")

    chunks = []
    for i in range(len(texts)):
        payload = {
            "filename": doc_filename(path), "doc_id": doc_id,
            "chunk_id": f"{doc_id}-page0-c{i}", "text": texts[i],
        }

        chunks.append(payload)

    meta = {
        "path": path, "doc_id": doc_id, "number_of_pages": number_of_pages,
        "chunk_size": 0, "chunk_overlap": 0, "number_of_chunks": len(chunks),
    }

    return {"chunks": chunks, "meta": meta}


def md2chunks(path, doc_id, chunk_size, chunk_overlap):
    with open(path, encoding="utf-8") as f:
        text = f.read()

    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    docs = splitter.create_documents([text])

    chunks = []
    for i in range(len(docs)):
        payload = {
            "filename": doc_filename(path), "doc_id": doc_id,
            "chunk_id": f"{doc_id}-page0-c{i}", "text": docs[i],
        }

        chunks.append(payload)

    meta = {
        "path": path, "doc_id": doc_id, "number_of_pages": 0,
        "chunk_size": chunk_size, "chunk_overlap": chunk_overlap,
        "number_of_chunks": len(chunks),
    }

    return {"chunks": chunks, "meta": meta}


def pptx2chunks(path, doc_id):
    prs = pptx.Presentation(path)
    chunks = []
    number_of_pages = len(prs.slides)

    for i, slide in enumerate(prs.slides):
        paragraphs = []
        page = i+1

        for shape in slide.shapes:
            if not shape.has_text_frame: continue
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if text: paragraphs.append(text)

        if len(paragraphs) == 0: continue

        payload = {
            "filename": doc_filename(path), "doc_id": doc_id,
            "chunk_id": f"{doc_id}-page{page}-c{i}", "text": "\n".join(paragraphs),
        }
        chunks.append(payload)

    meta = {
        "path": path, "doc_id": doc_id, "number_of_pages": number_of_pages,
        "chunk_size": 0, "chunk_overlap": 0, "number_of_chunks": len(chunks),
    }

    return { "chunks": chunks, "meta": meta }


def pdf2chunks(path, doc_id, chunk_size=1000, chunk_overlap=100):
    number_of_pages = 0
    chunks = []

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size, chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", "。", ".", "，", ",", " ", ""],
    )

    for page in PyPDFLoader(path).load_and_split():
        number_of_pages += 1
        texts = splitter.split_text(page.page_content)
        page = page.metadata['page']

        for i in range(len(texts)):
            payload = {
                "filename": doc_filename(path), "doc_id": doc_id,
                "chunk_id": f"{doc_id}-page{page}-c{i}", "text": texts[i],
            }

            chunks.append(payload)

    meta = {
        "path": path, "doc_id": doc_id, "number_of_pages": number_of_pages,
        "chunk_size": chunk_size, "chunk_overlap": chunk_overlap,
        "number_of_chunks": len(chunks),
    }

    return { "chunks": chunks, "meta": meta }


# gpt-3.5-turbo, gpt-4, text-davinci-003, cl100k_base
def count_tokens(text, model_name="gpt-4"):
    encoding = tiktoken.encoding_for_model(model_name)
    return len(encoding.encode(text))

def paragraphs_to_chunks(paragraphs, max_tokens):
    chunks, current, count = [], "", 0

    for p in paragraphs:
        count += count_tokens(p)
        current += p + "\n"

        if count >= max_tokens:
            chunks.append(current.strip())
            current, count = "", 0

    if current:
        chunks.append(current.strip())

    return chunks

def epub2chunks(path, doc_id, chunk_size=1000):
    def clean_text(text):
        text = re.sub(r'\s+', ' ', text)
        return text.strip()

    book, chunks = epub.read_epub(path), []

    for item in book.get_items():
        if item.get_type() == ebooklib.ITEM_DOCUMENT:
            soup = BeautifulSoup(item.get_content(), "xml")
            #tag = soup.find(['h1', 'h2', 'title'])
            #title = tag.get_text(strip=True) if title_tag else "Untitled Chapter"

            paragraphs = [
                clean_text(p.get_text())
                for p in soup.find_all("p") if p.get_text(strip=True)
            ]

            results = paragraphs_to_chunks(paragraphs, max_tokens=chunk_size)

            for i, chunk in enumerate(results):
                chunk = {
                    "filename": doc_filename(path), "doc_id": doc_id,
                    "chunk_id": f"{doc_id}-page0-c{i}", "text": chunk,
                }
                chunks.append(chunk)

    meta = {
        "path": path, "doc_id": doc_id, "number_of_pages": 0,
        "chunk_size": chunk_size, "chunk_overlap": 0,
        "number_of_chunks": len(chunks),
    }

    return { "chunks": chunks, "meta": meta }
