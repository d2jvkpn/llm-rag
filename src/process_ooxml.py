#!/usr/bin/env python3
import docx, pptx

from process_utils import doc_filename


def docx2chunks(path, doc_id):
    filename = doc_filename(path)
    doc = docx.Document(path)
    texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    chunks = []
    for i in range(len(texts)):
        payload = {
            "filename": filename, "doc_id": doc_id,
            "chunk_id": f"{doc_id}-page0-c{i}", "text": texts[i],
        }

        chunks.append(payload)

    meta = {
        "path": path, "doc_id": doc_id, "number_of_pages": 0,
        "chunk_size": 0, "chunk_overlap": 0, "number_of_chunks": len(chunks),
    }

    return { "meta": meta, "chunks": chunks }


def pptx2chunks(path, doc_id):
    filename = doc_filename(path)
    prs = pptx.Presentation(path)
    chunks = []

    for i, slide in enumerate(prs.slides):
        paragraphs = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if text:
                   paragraphs.append(text)

        if len(paragraphs) == 0: continue

        payload = {
            "filename": filename, "doc_id": doc_id,
            "chunk_id": f"{doc_id}-page{i+1}-c0", "text": "\n".join(paragraphs),
        }
        chunks.append(payload)

    meta = {
        "path": path, "doc_id": doc_id, "number_of_pages": len(prs.slides),
        "chunk_size": 0, "chunk_overlap": 0, "number_of_chunks": len(chunks),
    }

    return { "meta": meta, "chunks": chunks }
